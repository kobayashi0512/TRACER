"""Add the public TRACER repository URL to a manuscript availability statement."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


REPOSITORY_URL = "https://github.com/kobayashi0512/TRACER"
AVAILABILITY_PREFIX = "The TRACER implementation, frozen synthetic packets"
REPLACEMENT = (
    "The TRACER implementation, frozen synthetic packets, cohort registries, "
    "model protocols, machine-readable result summaries, reproducibility documents, "
    "and final editable figures are available in the public repository: "
)
AVAILABILITY_SUFFIX = (
    ". A data manifest records GEO source URLs and checksums; large public source "
    "files are provided as repository release assets. The real hard-negative analysis "
    "uses previously released, de-identified GEO study records [4]; reuse remains "
    "subject to the terms of each source repository."
)


def add_hyperlink(paragraph, url: str, text: str) -> None:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    run_properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)
    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def replace_availability_statement(document: Document) -> None:
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(AVAILABILITY_PREFIX):
            paragraph.clear()
            paragraph.add_run(REPLACEMENT)
            add_hyperlink(paragraph, REPOSITORY_URL, REPOSITORY_URL)
            paragraph.add_run(AVAILABILITY_SUFFIX)
            return
    raise ValueError("Data and Code Availability statement was not found.")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    document = Document(args.input)
    replace_availability_statement(document)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)


if __name__ == "__main__":
    main()
