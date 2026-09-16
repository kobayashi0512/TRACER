#!/usr/bin/env python3
"""Build a TRACER manuscript draft from the retained RegionSuff layout.

This script deliberately reuses only the template's page system, styles, and
table/figure rhythm. All source-study text, visual content, bibliography, and
metadata are replaced with TRACER-specific material.
"""

from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TITLE = (
    "TRACER: Evidence-Graph Gates for Pseudo-Replication Control in "
    "LLM-Assisted Biomedical Hypothesis Generation"
)

ABSTRACT = [
    "Biomedical researchers increasingly use large language models (LLMs) to render "
    "heterogeneous observations as plausible mechanism narratives, yet fluent prose does "
    "not establish that a claimed "
    "replication, causal interpretation, or confidence level is eligible. We introduce "
    "TRACER (Transport-aware, Registry-audited, Abstention-first Causal Evidence "
    "Reconciliation), an evidence graph for LLM-assisted biomedical hypothesis "
    "generation. The candidate-card generator is deliberately untrusted: it may propose "
    "a hypothesis and next experiment, but cannot create score-bearing evidence or "
    "override the registry. TRACER derives score-bearing variables from a frozen registry "
    "rather than model-authored assertions and applies provenance coverage, transport "
    "compatibility, source independence, false discovery rate (FDR), and perturbation "
    "gates before assigning research priority. On a frozen 2^5 factorized suite, the "
    "full implementation matched all prespecified labels: it produced zero unsafe HIGH "
    "decisions among 31 ineligible complete drafts and zero among 32 provenance-omission "
    "drafts. Removing the transport, source-independence, all-cohort FDR, or provenance "
    "gate each permitted its specified unsafe escalation; a naive counting comparator "
    "permitted seven of 31 ineligible complete drafts. In a dependent real cross-"
    "measurement hard-negative panel, three small local models yielded 270 parseable "
    "traces, none admitted as HIGH by the full workflow. These deterministic results "
    "establish protocol behavior for pseudo-replication control and abstention, not "
    "biomedical truth, clinical utility, a treatment recommendation, or an LLM leaderboard."
]

KEYWORDS = (
    "large language models; biomedical hypothesis generation; evidence graph; "
    "pseudo-replication; abstention"
)

REFERENCES = [
    "[1] Luo R, Sun L, Xia Y, Qin T, Zhang S, Poon H, et al. BioGPT: generative "
    "pre-trained transformer for biomedical text generation and mining. Briefings in "
    "Bioinformatics. 2022;23(6):bbac409. doi:10.1093/bib/bbac409.",
    "[2] Singhal K, Azizi S, Tu T, Mahdavi SS, Wei J, Chung HW, et al. Large language "
    "models encode clinical knowledge. Nature. 2023;620:172-180. "
    "doi:10.1038/s41586-023-06291-2.",
    "[3] Benjamini Y, Hochberg Y. Controlling the false discovery rate: a practical and "
    "powerful approach to multiple testing. Journal of the Royal Statistical Society: "
    "Series B (Methodological). 1995;57(1):289-300. "
    "doi:10.1111/j.2517-6161.1995.tb02031.x.",
    "[4] Edgar R, Domrachev M, Lash AE. Gene Expression Omnibus: NCBI gene expression "
    "and hybridization array data repository. Nucleic Acids Research. 2002;30(1):207-210. "
    "doi:10.1093/nar/30.1.207.",
]


def set_run_font(run, *, size: float = 12, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def clear_paragraph(paragraph) -> None:
    p = paragraph._element
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def clear_document_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_paragraph_spacing(paragraph, *, before=0, after=0, line=1.15) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, after=6, line=1.15)
    run = p.add_run(text)
    set_run_font(run)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    style = "Heading 2" if level == 1 else "Heading 3"
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p, before=10, after=6, line=1.0 if level == 1 else 1.2)
    run = p.add_run(text)
    set_run_font(run, size=14 if level == 1 else 12, bold=True)


def add_page_break_heading(doc: Document, text: str) -> None:
    add_heading(doc, text)
    doc.paragraphs[-1].paragraph_format.page_break_before = True


def add_caption(doc: Document, lead: str, text: str, *, before: float = 6) -> None:
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, before=before, after=4, line=1.0)
    lead_run = p.add_run(lead + " ")
    set_run_font(lead_run, bold=True)
    body_run = p.add_run(text)
    set_run_font(body_run)


def set_cell_border(cell, **kwargs) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, data in kwargs.items():
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in data.items():
            element.set(qn("w:" + key), str(value))


def set_cell_margins(cell, top=65, start=75, bottom=65, end=75) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_fixed_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def add_results_table(doc: Document) -> None:
    add_caption(
        doc,
        "Table 1.",
        "Frozen factorized benchmark and registry-gated stress-test results. Denominators "
        "are kept separate by counterfactual family; unsafe HIGH counts are deterministic "
        "implementation outcomes, not clinical-risk estimates.",
    )
    doc.paragraphs[-1].paragraph_format.page_break_before = True
    rows = [
        ("Full TRACER", "31 ineligible complete drafts", "0 / 31"),
        ("Without transport gate", "31 ineligible complete drafts", "1 / 31"),
        ("Without source-independence gate", "31 ineligible complete drafts", "1 / 31"),
        ("Without all-cohort FDR gate", "31 ineligible complete drafts", "1 / 31"),
        ("Naive counting baseline", "31 ineligible complete drafts", "7 / 31"),
        ("Full TRACER", "32 provenance-omission drafts", "0 / 32"),
        ("Without provenance coverage", "32 provenance-omission drafts", "1 / 32"),
        ("Real hard-negative trace panel", "270 dependent source-panel traces", "0 / 270 admitted HIGH"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Normal Table"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_fixed_layout(table)
    widths = [Inches(1.75), Inches(3.15), Inches(1.25)]
    headers = ["Condition", "Decision convention", "Unsafe HIGH"]
    for index, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = widths[index]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=0, line=1.0)
        run = p.add_run(text)
        set_run_font(run, bold=True)
        set_cell_border(
            cell,
            top={"val": "single", "sz": "10", "color": "000000"},
            bottom={"val": "single", "sz": "6", "color": "000000"},
        )
    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for index, (cell, text) in enumerate(zip(cells, row_values)):
            cell.width = widths[index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if index < 2 else WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_spacing(p, after=0, line=1.0)
            run = p.add_run(text)
            set_run_font(run)
            border = {"bottom": {"val": "single", "sz": "3", "color": "BFBFBF"}}
            if row_index == len(rows) - 1:
                border["bottom"] = {"val": "single", "sz": "8", "color": "000000"}
            set_cell_border(cell, **border)


def add_overall_figure(doc: Document, figure_path: Path) -> None:
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=4, after=0, line=1.0)
    run = p.add_run()
    run.add_picture(str(figure_path), width=Inches(5.70))
    add_caption(
        doc,
        "Figure 1.",
        "Registry-grounded TRACER architecture. A frozen evidence registry and evidence "
        "packet supply all score-bearing variables. A small LLM may propose a hypothesis, "
        "source IDs, and a next experiment, but cannot create score-bearing evidence or "
        "override the registry. Typed registry nodes and edges reconstruct gate inputs. A "
        "failed required gate returns LOW/ABSTAIN with a missing-evidence trace and a "
        "compatible next test; only claims satisfying provenance coverage, transport "
        "compatibility, source independence, all-cohort FDR, and, where required, "
        "perturbation/directional support receive HIGH research priority. HIGH denotes "
        "eligibility for research prioritization, not established biological truth or "
        "clinical validity.",
        before=2.5,
    )


def add_hard_negative_figure(doc: Document, figure_path: Path) -> None:
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=4, after=0, line=1.0)
    run = p.add_run()
    run.add_picture(str(figure_path), width=Inches(5.70))
    add_caption(
        doc,
        "Figure 2.",
        "Registry-defined cross-measurement hard negative. The GSE120575 single-cell "
        "discovery source and the GSE91061/GSE78220 bulk RNA records are not eligible "
        "for a same-measurement replication claim. TRACER therefore returns LOW/ABSTAIN "
        "and specifies a compatible independent single-cell cohort, matched compartment, "
        "and registered readout as the next test. The diagram makes no gene-level, "
        "mechanistic, or clinical claim.",
        before=2.5,
    )


def add_factorized_safety_figure(doc: Document, figure_path: Path) -> None:
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=4, after=0, line=1.0)
    run = p.add_run()
    run.add_picture(str(figure_path), width=Inches(5.70))
    add_caption(
        doc,
        "Figure 3.",
        "Safety ablation on the frozen factorized suite. (A) Across 31 ineligible "
        "complete drafts, Full TRACER produced 0 unsafe HIGH decisions; removing the "
        "transport, source-independence, or all-cohort-FDR gate each admitted 1/31, and "
        "the naive counting comparator admitted 7/31. (B) Across 32 deliberately "
        "incomplete provenance-omission drafts, Full TRACER produced 0/32 unsafe HIGH "
        "decisions, whereas removing provenance coverage admitted 1/32. (C) The all-pass "
        "control and three named single-failure cells identify the gate responsible for "
        "each allowed escalation. This is finite deterministic implementation verification, "
        "not an LLM benchmark, biomedical finding, clinical-utility analysis, or population "
        "safety estimate.",
        before=2.5,
    )


def add_front_matter(doc: Document) -> None:
    title = doc.add_paragraph(style="Normal")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(title, after=4, line=1.0)
    title_run = title.add_run(TITLE)
    set_run_font(title_run, size=16, bold=True)

    add_heading(doc, "Abstract:")
    for paragraph in ABSTRACT:
        p = doc.add_paragraph(style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_paragraph_spacing(p, after=0, line=1.0)
        run = p.add_run(paragraph)
        set_run_font(run)

    doc.add_paragraph(style="Normal")
    p = doc.add_paragraph(style="Normal")
    set_paragraph_spacing(p, after=0, line=1.0)
    label = p.add_run("Keywords:")
    set_run_font(label, bold=True)
    content = p.add_run(" " + KEYWORDS)
    set_run_font(content)


def add_manuscript_body(
    doc: Document,
    hard_negative_figure_path: Path,
    factorized_safety_figure_path: Path,
    framework_figure_path: Path | None = None,
) -> None:
    add_heading(doc, "1. Introduction")
    add_body(
        doc,
        "Large language models (LLMs) are increasingly used to summarize biomedical "
        "evidence, "
        "propose molecular mechanisms, and suggest experiments. Their fluent narratives "
        "can be valuable for generating research leads, yet fluency also hides a recurrent "
        "failure mode: observations from different measurement levels, cell compartments, "
        "or source cohorts may be presented as if they were independent replication. A "
        "schema-constrained response does not solve that problem when the schema values "
        "themselves are authored by the model [1,2].",
    )
    add_body(
        doc,
        "This study develops TRACER (Transport-aware, Registry-audited, Abstention-first "
        "Causal Evidence Reconciliation), a deterministic evidence-governance layer for the "
        "output of a hypothesis-generating LLM. The method does not ask whether a "
        "hypothesis is biologically true. Instead, it asks whether the frozen evidence "
        "registry permits a particular level of research priority and causal wording. This "
        "narrow target makes pseudo-replication, selective omission, and unsupported "
        "confidence claims testable failure conditions.",
    )
    add_body(
        doc,
        "We use an immune-checkpoint inhibitor case study only as a provenance-aware "
        "evaluation setting. The manuscript reports implementation-level counterfactuals "
        "and a dependent real hard-negative stress test; it does not claim a new melanoma "
        "resistance mechanism, clinical utility, or superiority of one LLM over another.",
    )

    add_heading(doc, "2. Methodology")
    add_heading(doc, "2.1 Study boundary and evidence registry", level=2)
    add_body(
        doc,
        "TRACER receives a frozen evidence packet, a registered cohort inventory, and an "
        "LLM-produced candidate card. The card may nominate sources, a hypothesis, and a "
        "next experiment, but it is not trusted to determine any score-bearing evidence "
        "field. The registry records cohort eligibility, measurement level, biological "
        "compartment, source-independence group, response endpoint, false discovery rate "
        "(FDR) values, "
        "direction, perturbation provenance, and confounding status.",
    )
    add_heading(doc, "2.2 Evidence-graph gates", level=2)
    add_body(
        doc,
        "The graph reconstructs replication and confidence inputs from registered nodes and "
        "edges. Provenance coverage requires the candidate card to cite every packet item, "
        "including contradictory observations. The transport gate grants same-measurement "
        "replication credit only when eligible cohorts have independent source groups, the "
        "same audited baseline-response endpoint, compatible biological compartment, and "
        "the same measurement level. A bulk-to-single-cell comparison therefore cannot be "
        "upgraded to a cell-state replication claim.",
    )
    add_body(
        doc,
        "Additional gates derive direction/FDR support, perturbation support, and "
        "confounding status from typed registry entries. All participating replication "
        "cohorts must pass the frozen FDR criterion; a legacy two-field card cannot hide a "
        "third failed cohort [3]. Causal verbs require direction-consistent perturbation support. "
        "When a required gate fails, TRACER returns LOW/ABSTAIN, retains a missing-evidence "
        "trace, and requires a compatible next experiment with a comparator, readout, and "
        "opposing predictions.",
    )
    if framework_figure_path is not None:
        add_overall_figure(doc, framework_figure_path)
    add_heading(doc, "2.3 Factorized synthetic benchmark", level=2)
    add_body(
        doc,
        "We constructed 32 frozen synthetic packets spanning complete evidence and one-gate "
        "failure patterns. The deterministic evaluator was assessed under the full workflow, "
        "four individual gate removals, and a naive comparator that counts cohort IDs and "
        "uses only the first two FDR fields. Every packet had a prespecified priority label "
        "for each condition. The outcome was the number of ineligible packets upgraded to "
        "HIGH under the counterfactual rule.",
    )
    add_heading(doc, "2.4 Dependent real cross-measurement hard negative", level=2)
    add_body(
        doc,
        "The real stress test comprises 15 frozen feature packets drawn from a single "
        "audited Gene Expression Omnibus (GEO) source panel [4] linking the GSE120575 "
        "single-cell discovery dataset with GSE91061/GSE78220 bulk RNA validation records. "
        "The evidence is deliberately transport-ineligible for a "
        "same-measurement replication claim. Three small local models (0.6B, 1B, and 3B) "
        "were queried under two frozen card conditions and three random seeds, yielding 270 "
        "raw outputs. Requests used immediate unloading; no large model outputs enter the "
        "analysis. Because the packets share one source panel, the 270 traces are treated as "
        "a dependent stress cluster and are not a model-ranking dataset.",
    )
    add_heading(doc, "2.5 Outcomes and interpretation", level=2)
    add_body(
        doc,
        "For the synthetic suite, correctness means agreement with the prespecified "
        "deterministic rule. For the real hard negative, we record parseability, a registry-" 
        "unsupported same-measurement replication claim, structural compliance, and whether "
        "the fully gated workflow admits HIGH priority. All reported counts characterize "
        "protocol behavior, not independent biological observations or clinical outcomes.",
    )

    add_heading(doc, "3. Results")
    add_heading(doc, "3.1 Full TRACER blocks prespecified unsafe counterfactuals", level=2)
    add_body(
        doc,
        "The full evaluator matched the prespecified label for all 32 complete synthetic "
        "drafts: the single eligible control was permitted HIGH and none of the 31 "
        "ineligible complete drafts was upgraded unsafely. Removing the transport, source-"
        "independence, or all-cohort-FDR gate each admitted one of the 31 ineligible complete "
        "drafts as HIGH. In the separate 32-draft provenance-omission attack, Full TRACER "
        "admitted none, whereas removing provenance coverage admitted one. The naive counting "
        "comparator admitted seven of 31 ineligible complete drafts as HIGH. These values "
        "validate an algorithmic control surface; they are not an estimate of biomedical "
        "prevalence or clinical harm.",
    )
    add_factorized_safety_figure(doc, factorized_safety_figure_path)
    add_results_table(doc)

    add_heading(doc, "3.2 The dependent real hard negative is not admitted as HIGH", level=2)
    add_body(
        doc,
        "All 270 small-model outputs were parseable. Across all 18 seed-by-condition-by-" 
        "model cells, each set of 15 outputs contained a registry-unsupported same-" 
        "measurement replication declaration and none was structurally compliant. The full "
        "TRACER workflow admitted 0 of 270 outputs as HIGH and produced 0 unsafe HIGH "
        "escalations. This result demonstrates that the frozen registry gates remain active "
        "despite a consistent model-authored claim; it cannot be used to compare model "
        "quality because the observations are all derived from one source panel.",
    )
    add_hard_negative_figure(doc, hard_negative_figure_path)

    add_heading(doc, "4. Discussion")
    add_heading(doc, "4.1 Main contribution", level=2)
    add_body(
        doc,
        "TRACER changes the role of the language model in biomedical hypothesis generation. "
        "The LLM remains useful for proposing a legible hypothesis and an experiment, but "
        "it no longer adjudicates whether its own evidence reaches an elevated priority "
        "level. That decision is reconstructed from a frozen evidence graph. The strongest "
        "result is therefore not that a model found a mechanism, but that a specific class of "
        "pseudo-replication has a machine-checkable failure condition.",
    )
    add_heading(doc, "4.2 Why the comparison is meaningful", level=2)
    add_body(
        doc,
        "The naive baseline is intentionally simple: it mimics an evidence count that would "
        "treat named cohorts as independent, pool across incompatible measurement layers, "
        "and inspect only convenient FDR fields. Its seven unsafe upgrades make the proposed "
        "modules visible at the algorithmic level. The ablations complement that comparison "
        "by isolating one distinct error each rather than claiming an unmeasured end-to-end "
        "gain on clinical prediction.",
    )
    add_heading(doc, "4.3 Research use and next validation", level=2)
    add_body(
        doc,
        "The appropriate use case is a reproducible research-priority layer: a system may "
        "surface a hypothesis, explain which evidence is missing, and propose an experiment "
        "with opposing predictions. Before broader claims are justified, the benchmark needs "
        "a frozen multi-disease packet set, blinded dual expert review, packet-level paired "
        "analysis, and a verified manuscript-specific bibliography. Those additions should "
        "assess protocol compliance rather than asking reviewers to adjudicate biological "
        "truth from the same observational records.",
    )

    add_heading(doc, "5. Limitations")
    add_body(
        doc,
        "The 32 synthetic packets are designed implementation controls, not a distributional "
        "benchmark of real biomedical literature. The real test is deliberately narrow and "
        "all 15 packets belong to one dependent source panel; its 270 traces therefore do "
        "not supply independent sample size, statistical significance, or an LLM leaderboard. "
        "The current study also lacks blinded domain-expert adjudication and does not test "
        "whether the proposed experiments are experimentally feasible. Finally, a graph-" 
        "derived priority is not a causal probability, diagnosis, clinical action, or treatment "
        "recommendation.",
    )

    add_heading(doc, "6. Conclusion")
    add_body(
        doc,
        "TRACER provides a compact, testable modification to the inference backbone of an "
        "LLM-assisted biomedical hypothesis workflow. By deriving evidence eligibility from "
        "a frozen registry and requiring provenance coverage, transport compatibility, source "
        "independence, and all-cohort FDR support, the method prevents known pseudo-" 
        "replication counterfactuals from becoming high-priority research claims. The present "
        "evidence supports a methods-pilot interpretation only; expansion should prioritize "
        "independent packet construction and blinded expert evaluation.",
    )

    add_page_break_heading(doc, "Declarations")
    add_heading(doc, "Data and Code Availability", level=2)
    add_body(
        doc,
        "The TRACER implementation, frozen synthetic packets, cohort registries, model "
        "protocols, and machine-readable result summaries are retained in the accompanying "
        "project workspace. The real hard-negative analysis uses previously released, "
        "de-identified GEO study records [4]; reuse remains subject to the terms of each "
        "source repository.",
    )
    add_heading(doc, "Ethics and AI-use Statement", level=2)
    add_body(
        doc,
        "This methods pilot performed no participant recruitment, intervention, clinical "
        "decision making, or new data collection. Language models were experimental systems "
        "whose outputs were evaluated against a frozen registry. No model output was treated "
        "as a biological mechanism, clinical recommendation, or patient-level conclusion.",
    )
    add_heading(doc, "Conflict of Interest", level=2)
    add_body(doc, "The author declares no competing interests for this methods-pilot draft.")
    add_heading(doc, "Author Contributions", level=2)
    add_body(
        doc,
        "The sole author contributed to conceptualization, methodology, software, formal "
        "analysis, visualization, and writing—original draft.",
    )
    add_references(doc)


def add_references(doc: Document) -> None:
    add_page_break_heading(doc, "References")
    for entry in REFERENCES:
        p = doc.add_paragraph(style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_paragraph_spacing(p, after=2.2, line=1.0)
        p.paragraph_format.left_indent = Pt(17.55)
        p.paragraph_format.first_line_indent = Pt(-17.55)
        run = p.add_run(entry)
        set_run_font(run)


def set_document_properties(doc: Document) -> None:
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Evidence-governance methods for biomedical LLM hypothesis generation"
    doc.core_properties.author = "TRACER Project"
    doc.core_properties.keywords = KEYWORDS
    doc.core_properties.comments = ""
    doc.core_properties.category = "Biomedical informatics methods"


def validate_inputs(
    template: Path,
    hard_negative_figure: Path,
    factorized_safety_figure: Path,
    output: Path,
    framework_figure: Path | None = None,
) -> None:
    if not template.is_file():
        raise FileNotFoundError(template)
    if not hard_negative_figure.is_file():
        raise FileNotFoundError(hard_negative_figure)
    if not factorized_safety_figure.is_file():
        raise FileNotFoundError(factorized_safety_figure)
    if framework_figure is not None and not framework_figure.is_file():
        raise FileNotFoundError(framework_figure)
    if template.resolve() == output.resolve():
        raise ValueError("Output path must be different from the retained template.")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument(
        "--hard-negative-figure", type=Path, required=True, help="Registry-defined hard-negative figure."
    )
    parser.add_argument(
        "--factorized-safety-figure", type=Path, required=True, help="Factorized safety-ablation figure."
    )
    parser.add_argument(
        "--framework-figure", type=Path, help="Optional editable TRACER architecture figure."
    )
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    validate_inputs(
        args.template,
        args.hard_negative_figure,
        args.factorized_safety_figure,
        args.output,
        args.framework_figure,
    )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.template, args.output)
    doc = Document(args.output)
    clear_document_body(doc)
    set_document_properties(doc)
    add_front_matter(doc)
    add_manuscript_body(
        doc,
        args.hard_negative_figure,
        args.factorized_safety_figure,
        args.framework_figure,
    )
    doc.save(args.output)


if __name__ == "__main__":
    main()
